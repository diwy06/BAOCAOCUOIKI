#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ghép nguồn Markdown trong docs/bao_cao/ thành tệp DOCX báo cáo.

Cách dùng (từ thư mục gốc repo):
    python docs/build_bao_cao.py                 # bản đầy đủ
    python docs/build_bao_cao.py --den 20        # bản cắt đến hết Chương 2

Cả hai bản cắt từ cùng một nguồn nên phần đầu giống nhau từng trang.

Cú pháp Markdown:
    # H1 / ## H2 / ### H3      tiêu đề chương, mục, tiểu mục
    #! Tiêu đề                 tiêu đề phụ (MỤC LỤC...), không vào mục lục
    | a | b |                   bảng, dòng đầu là tiêu đề cột
    ``` ... ```                 khối mã
    - mục                       gạch đầu dòng
    > ghi chú                   đoạn ghi chú in nghiêng
    Bảng 1.1 ...                chú thích bảng (style CaptionBang)
    Hình 4.1 ...                chú thích hình (style CaptionHinh)
    ![...](đường dẫn)           chèn ảnh, phần trong [] là chú thích hình
    [[PAGEBREAK]]               ngắt trang
    [[SECT]]                    sang section mới, số trang Ả Rập từ 1
    [[TOC]] / [[LOT]] / [[LOF]] mục lục / danh sách bảng / danh sách hình
    ::: bia ... :::             trang bìa
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "docs" / "bao_cao"
DOCS_DIR = ROOT / "docs"

FONT_BODY = "Times New Roman"
FONT_CODE = "Consolas"
SIZE_BODY = 13
SIZE_TABLE = 10.5
SIZE_CAPTION = 11

HEADER_TEXT = "Báo cáo đồ án INT13185 — Trần Hữu Duy — N24DCDT022 — DT011"

STYLE_CAP_BANG = "CaptionBang"
STYLE_CAP_HINH = "CaptionHinh"
STYLE_TIEU_DE = "TieuDePhu"

SECTPR_AFTER_PGNUM = (
    "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
    "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
    "w:printerSettings", "w:sectPrChange",
)


# --------------------------------------------------------------------------
# Tiện ích cấp thấp
# --------------------------------------------------------------------------
def set_run_font(run, name=FONT_BODY, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(element, fill):
    pr = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def add_field(paragraph, instruction, placeholder=""):
    """Chèn field Word (PAGE, TOC...).

    Mỗi thành phần của field phải nằm trong một run riêng; gộp tất cả vào một
    run sẽ tạo field hỏng, Word có thể treo khi cập nhật.
    """
    def run_with(child):
        r = paragraph.add_run()
        r._r.append(child)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run_with(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run_with(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run_with(separate)

    if placeholder:
        paragraph.add_run(placeholder)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def set_page_numbering(section, fmt="decimal", start=1):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:fmt"), fmt)
    el.set(qn("w:start"), str(start))
    sectPr.insert_element_before(el, *SECTPR_AFTER_PGNUM)


# --------------------------------------------------------------------------
# Định dạng nội tuyến
# --------------------------------------------------------------------------
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_inline(paragraph, text, size=SIZE_BODY, bold=False, italic=False, name=FONT_BODY):
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            run = paragraph.add_run(piece[2:-2])
            set_run_font(run, name, size, bold=True, italic=italic)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, FONT_CODE, size - 1.5, bold=bold, italic=False)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, name, size, bold=bold, italic=True)
        else:
            run = paragraph.add_run(piece)
            set_run_font(run, name, size, bold=bold, italic=italic)


# --------------------------------------------------------------------------
# Document và kiểu
# --------------------------------------------------------------------------
def setup_document():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True
    set_page_numbering(section, "lowerRoman", 1)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(SIZE_BODY)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    pf = normal.paragraph_format
    pf.line_spacing = 1.25
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    specs = {
        "Heading 1": (16, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 14),
        "Heading 2": (14, True, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
        "Heading 3": (13, True, WD_ALIGN_PARAGRAPH.LEFT, 10, 4),
    }
    for name, (size, bold, align, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT_BODY
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.2
        st.paragraph_format.keep_with_next = True

    for style_name in (STYLE_CAP_BANG, STYLE_CAP_HINH):
        st = doc.styles.add_style(style_name, 1)  # 1 = paragraph style
        st.base_style = doc.styles["Normal"]
        st.font.name = FONT_BODY
        st.font.size = Pt(SIZE_CAPTION)
        st.font.bold = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.keep_with_next = True
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

    st = doc.styles.add_style(STYLE_TIEU_DE, 1)
    st.base_style = doc.styles["Normal"]
    st.font.name = FONT_BODY
    st.font.size = Pt(16)
    st.font.bold = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.first_line_indent = Cm(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(14)
    st.paragraph_format.page_break_before = True
    st.paragraph_format.keep_with_next = True
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(HEADER_TEXT)
    set_run_font(hr, FONT_BODY, 9, italic=True, color=RGBColor(0x55, 0x55, 0x55))

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("Trang ")
    set_run_font(r1, FONT_BODY, 10)
    add_field(fp, "PAGE", "1")

    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)

    return doc


def add_body_section(doc):
    """Sang section mới: đánh số trang Ả Rập bắt đầu từ 1."""
    new = doc.add_section(WD_SECTION.NEW_PAGE)
    new.top_margin = Cm(2.5)
    new.bottom_margin = Cm(2.5)
    new.left_margin = Cm(3.0)
    new.right_margin = Cm(2.0)
    new.different_first_page_header_footer = False
    set_page_numbering(new, "decimal", 1)
    return new


# --------------------------------------------------------------------------
# Các khối nội dung
# --------------------------------------------------------------------------
def add_heading(doc, text, level, first_block=False, suppress_break=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1 and not first_block and not suppress_break:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text.strip())
    set_run_font(run, FONT_BODY, {1: 16, 2: 14, 3: 13}[level], bold=True)
    return p


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.0)
    add_inline(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(1.0)
    pf.first_line_indent = Cm(-0.4)
    pf.space_after = Pt(3)
    add_inline(p, "• " + text)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.8)
    pf.right_indent = Cm(0.5)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    add_inline(p, text, size=SIZE_BODY - 0.5, italic=True)
    return p


def add_caption(doc, text, style=STYLE_CAP_BANG):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text.strip())
    set_run_font(run, FONT_BODY, SIZE_CAPTION, bold=True)
    return p


def add_code(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shade(p._p, "F4F4F4")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, FONT_CODE, 10)
        if i < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)
    return p


def add_table(doc, rows):
    n_cols = max(len(r) for r in rows)
    rows = [r + [""] * (n_cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    usable = 16.0  # cm
    weights = []
    for c in range(n_cols):
        longest = max(len(r[c]) for r in rows)
        weights.append(min(max(longest, 10), 46) + 6.0)
    total = sum(weights)
    widths = [max(usable * w / total, usable * 0.11) for w in weights]
    scale = usable / sum(widths)
    widths = [w * scale for w in widths]

    for c, w in enumerate(widths):
        table.columns[c].width = Cm(w)

    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Cm(widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.05
            pf.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or len(text) <= 8)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            if r_idx < len(rows) - 1:
                pf.keep_with_next = True
            add_inline(p, text, size=SIZE_TABLE, bold=(r_idx == 0))
            if r_idx == 0:
                shade(cell._tc, "E7E6E6")

    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))

    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_image(doc, path, caption):
    full = (ROOT / path).resolve()
    if not full.exists():
        p = doc.add_paragraph()
        add_inline(p, f"[thiếu ảnh: {path}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(full), width=Cm(15.2))
    if caption:
        add_caption(doc, caption, STYLE_CAP_HINH)


def add_cover(doc, lines):
    for _ in range(2):
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
    for raw in lines:
        if not raw.strip():
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        size, bold, text = 13.0, True, raw
        m = re.match(r"^\[([^\]]*)\]\s*(.*)$", raw)
        if m:
            opts, text = m.group(1), m.group(2)
            for opt in opts.split(","):
                opt = opt.strip()
                if opt.startswith("s"):
                    size = float(opt[1:])
                elif opt == "b":
                    bold = True
                elif opt == "n":
                    bold = False
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        set_run_font(run, FONT_BODY, size, bold=bold)


def add_toc_field(doc, instruction, placeholder):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_field(p, instruction, placeholder)
    return p


# --------------------------------------------------------------------------
# Bộ phân tích Markdown
# --------------------------------------------------------------------------
def parse_blocks(text):
    blocks = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "::: bia":
            i += 1
            buf = []
            while i < len(lines) and lines[i].strip() != ":::":
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("cover", buf))
            continue

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", buf))
            continue

        if stripped == "[[PAGEBREAK]]":
            blocks.append(("pagebreak", None))
            i += 1
            continue

        if stripped == "[[SECT]]":
            blocks.append(("sect", None))
            i += 1
            continue

        if stripped in ("[[TOC]]", "[[LOT]]", "[[LOF]]"):
            blocks.append((stripped.strip("[]").lower(), None))
            i += 1
            continue

        if stripped.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if not re.fullmatch(r"\|[\s:\-|]+\|", row):
                    row = row.replace("\\|", "\x01")
                    cells = [c.replace("\x01", "|").strip() for c in row.strip("|").split("|")]
                    buf.append(cells)
                i += 1
            blocks.append(("table", buf))
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2)))
            i += 1
            continue

        if stripped.startswith("#!"):
            blocks.append(("title", stripped[2:].strip()))
            i += 1
            continue

        if stripped.startswith("> "):
            blocks.append(("note", stripped[2:]))
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            blocks.append(("bullet", re.sub(r"^[-*]\s+", "", stripped)))
            i += 1
            continue

        img = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if img:
            blocks.append(("image", (img.group(2), img.group(1))))
            i += 1
            continue

        m = re.match(r"^(Bảng|Hình)\s+\S+", stripped)
        if m:
            kind = STYLE_CAP_BANG if m.group(1) == "Bảng" else STYLE_CAP_HINH
            blocks.append(("caption", (stripped, kind)))
            i += 1
            continue

        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("|", "```", "#", "> ", "- ", "* "))
                    or re.match(r"^!\[", nxt)
                    or nxt in ("[[PAGEBREAK]]", "[[TOC]]", "[[LOT]]", "[[LOF]]", "[[SECT]]", "::: bia")):
                break
            buf.append(nxt)
            i += 1
        blocks.append(("p", " ".join(buf)))

    return blocks


def select_sources(den=None, tu=None):
    sources = sorted(SRC_DIR.glob("*.md"))
    if den is not None:
        sources = [s for s in sources if int(s.name[:2]) <= den]
    if tu is not None:
        sources = [s for s in sources if int(s.name[:2]) >= tu]
    return sources


def build(den=None, tu=None, out=None):
    if not SRC_DIR.exists():
        sys.exit(f"Không thấy thư mục nguồn: {SRC_DIR}")

    sources = select_sources(den, tu)
    if not sources:
        sys.exit(f"Không có tệp .md nào trong {SRC_DIR}")

    out_path = Path(out) if out else DOCS_DIR / "N24DCDT022_DT011_TranHuuDuy_BaoCao.docx"

    doc = setup_document()
    if tu is not None:
        # phần sau (bắt đầu từ Chương 3): số trang Ả Rập từ 1, có số ngay trang đầu
        doc.sections[0].different_first_page_header_footer = False
        set_page_numbering(doc.sections[0], "decimal", 1)
    first_block = True
    vua_sang_section = False

    for src in sources:
        for kind, payload in parse_blocks(src.read_text(encoding="utf-8-sig")):
            if kind == "cover":
                add_cover(doc, payload)
            elif kind in ("h1", "h2", "h3"):
                add_heading(doc, payload, int(kind[1]), first_block,
                            suppress_break=vua_sang_section)
                vua_sang_section = False
            elif kind == "title":
                tp = doc.add_paragraph(style=STYLE_TIEU_DE)
                trun = tp.add_run(payload)
                set_run_font(trun, FONT_BODY, 16, bold=True)
            elif kind == "p":
                add_paragraph(doc, payload)
            elif kind == "bullet":
                add_bullet(doc, payload)
            elif kind == "note":
                add_note(doc, payload)
            elif kind == "caption":
                add_caption(doc, payload[0], payload[1])
            elif kind == "code":
                add_code(doc, payload)
            elif kind == "table":
                add_table(doc, payload)
            elif kind == "image":
                add_image(doc, payload[0], payload[1])
            elif kind == "pagebreak":
                doc.add_page_break()
            elif kind == "sect":
                add_body_section(doc)
                vua_sang_section = True
            elif kind == "toc":
                add_toc_field(doc, 'TOC \\o "1-3" \\h \\z \\u',
                              "Mục lục sẽ được cập nhật khi mở bằng Word (Ctrl+A, F9).")
            elif kind == "lot":
                add_toc_field(doc, f'TOC \\h \\z \\t "{STYLE_CAP_BANG},1"',
                              "Danh sách bảng sẽ được cập nhật khi mở bằng Word (Ctrl+A, F9).")
            elif kind == "lof":
                add_toc_field(doc, f'TOC \\h \\z \\t "{STYLE_CAP_HINH},1"',
                              "Danh sách hình sẽ được cập nhật khi mở bằng Word (Ctrl+A, F9).")
            first_block = False

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Đã xuất: {out_path}")
    print(f"  nguồn: {len(sources)} tệp .md ({sources[0].name} .. {sources[-1].name})")
    print(f"  số đoạn: {len(doc.paragraphs)}, số bảng: {len(doc.tables)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Ghép nguồn Markdown thành DOCX báo cáo.")
    ap.add_argument("--den", type=int, default=None,
                    help="chỉ lấy các tệp nguồn có số thứ tự <= DEN")
    ap.add_argument("--tu", type=int, default=None,
                    help="chỉ lấy các tệp nguồn có số thứ tự >= TU")
    ap.add_argument("--out", default=None, help="đường dẫn tệp DOCX đích")
    a = ap.parse_args()
    build(a.den, a.tu, a.out)
